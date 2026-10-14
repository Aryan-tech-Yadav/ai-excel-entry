"""
Core logic ke tests — internet ya API key ki zaroorat nahi.

Chalane ke liye (project folder se):
    python -m unittest discover -s tests -v
"""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import config
from core import ai_parser, excel_store, validator


class ValidatorTests(unittest.TestCase):
    def test_amount_formats(self):
        self.assertEqual(validator.parse_amount("Rs. 5,000/-"), 5000.0)
        self.assertEqual(validator.parse_amount("₹1,23,456.50"), 123456.5)
        self.assertEqual(validator.parse_amount(750), 750.0)
        self.assertIsNone(validator.parse_amount("paanch hazaar"))
        self.assertIsNone(validator.parse_amount(None))

    def test_date_formats(self):
        for text in ["24-09-2026", "24/09/2026", "2026-09-24", "24 Sep 2026", "24.09.2026"]:
            self.assertEqual(validator.parse_date(text).isoformat(), "2026-09-24", text)
        self.assertIsNone(validator.parse_date("kal shaam"))

    def test_validate_record_ok(self):
        clean, issues = validator.validate_record(
            {"Name": " Ramesh ", "Amount": "₹5,000", "Date": "2026-09-24"})
        self.assertEqual(issues, [])
        self.assertEqual(clean["Name"], "Ramesh")
        self.assertEqual(clean["Amount"], 5000.0)
        self.assertEqual(clean["Date"], "24-09-2026")

    def test_validate_record_problems(self):
        _, issues = validator.validate_record({"Name": None, "Amount": "abc", "Date": "??"})
        joined = " ".join(issues)
        self.assertIn("Name", joined)
        self.assertIn("Amount", joined)
        self.assertIn("Date", joined)

    def test_entry_id_stable(self):
        a = {"Name": "Ramesh", "Amount": 5000.0, "Date": "24-09-2026", "Invoice_No": "INV-1"}
        b = {"Name": "ramesh ", "Amount": 5000.0, "Date": "24-09-2026", "Invoice_No": "inv-1"}
        self.assertEqual(validator.make_entry_id(a), validator.make_entry_id(b))


class ExcelStoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self._old = (config.EXCEL_PATH, config.BACKUP_DIR)
        config.EXCEL_PATH = root / "test.xlsx"
        config.BACKUP_DIR = root / "backups"

    def tearDown(self):
        config.EXCEL_PATH, config.BACKUP_DIR = self._old
        self.tmp.cleanup()

    def _rec(self, name="Ramesh", amount=5000.0, inv="INV-1"):
        return {"Name": name, "Amount": amount, "Date": "24-09-2026",
                "Invoice_No": inv, "Notes": None}

    def test_append_and_duplicate(self):
        r1 = excel_store.append_records([self._rec()], source="test")
        self.assertEqual(r1, {"added": 1, "duplicates": 0})
        r2 = excel_store.append_records([self._rec(), self._rec("Sita", 200.0, "INV-2")])
        self.assertEqual(r2, {"added": 1, "duplicates": 1})
        df = excel_store.read_all()
        self.assertEqual(len(df), 2)
        self.assertEqual(list(df["Name"]), ["Ramesh", "Sita"])

    def test_backup_created(self):
        excel_store.append_records([self._rec()])
        excel_store.append_records([self._rec("Sita", 200.0, "INV-2")])
        self.assertTrue(list(config.BACKUP_DIR.glob("*.xlsx")))

    def test_new_field_adds_column_to_old_file(self):
        excel_store.append_records([self._rec()])
        old_fields = config.FIELDS
        config.FIELDS = old_fields + [config.Field("City", "text", "City name")]
        try:
            excel_store.append_records([{**self._rec("Sita", 200.0, "INV-2"), "City": "Pune"}])
            ws = load_workbook(config.EXCEL_PATH).active
            header = [c.value for c in ws[1]]
            self.assertIn("City", header)
            self.assertEqual(ws.cell(row=3, column=header.index("City") + 1).value, "Pune")
        finally:
            config.FIELDS = old_fields

    def test_existing_ids(self):
        excel_store.append_records([self._rec()])
        self.assertEqual(len(excel_store.existing_ids()), 1)

    def test_stats(self):
        excel_store.append_records([self._rec(amount=100.0), self._rec("Sita", 250.5, "INV-2")])
        s = excel_store.stats()
        self.assertEqual(s["count"], 2)
        self.assertAlmostEqual(s["total"], 350.5)


class AIReplyParsingTests(unittest.TestCase):
    def test_plain_json(self):
        out = ai_parser._extract_records('{"records": [{"Name": "A", "Amount": 1}]}')
        self.assertEqual(out, [{"Name": "A", "Amount": 1}])

    def test_code_fence_and_chatter(self):
        reply = 'Sure! Here you go:\n```json\n{"records": [{"Name": "B"}]}\n```'
        self.assertEqual(ai_parser._extract_records(reply), [{"Name": "B"}])

    def test_single_dict_and_list(self):
        self.assertEqual(ai_parser._extract_records('{"Name": "C"}'), [{"Name": "C"}])
        self.assertEqual(ai_parser._extract_records('[{"Name": "D"}]'), [{"Name": "D"}])

    def test_garbage_raises(self):
        with self.assertRaises(ai_parser.AIParseError):
            ai_parser._extract_records("no json here")


if __name__ == "__main__":
    unittest.main()
