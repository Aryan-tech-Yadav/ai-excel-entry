"""
excel_store.py — Excel file ke saath saara kaam yahin hota hai.

Safety features:
  * har save se pehle automatic backup (backups/ folder)
  * duplicate entry skip
  * file Excel mein khuli ho toh crash nahi, saaf message
  * config mein naya field add karo toh Excel mein naya column apne-aap jud jata hai
  * temp file mein save karke replace — beech mein crash hua toh original file safe rehti hai
"""
import os
import shutil
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from core import validator

log = config.get_logger("excel_store")

SYSTEM_COLS = ["Source", "Added_On", "Entry_ID"]
_HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", start_color="1F4E78")
_BODY_FONT = Font(name="Arial")


class ExcelLockedError(Exception):
    """Excel file kisi aur program (aksar Excel) mein khuli hai."""


def headers() -> list:
    return [f.name for f in config.FIELDS] + SYSTEM_COLS


# ---------------------------------------------------------------- internal helpers
def _style_header_cell(cell):
    cell.font = _HEADER_FONT
    cell.fill = _HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _get_sheet(wb):
    if config.SHEET_NAME in wb.sheetnames:
        return wb[config.SHEET_NAME]
    ws = wb.create_sheet(config.SHEET_NAME)
    return ws


def _ensure_workbook():
    """File na ho toh header ke saath nayi bana deta hai."""
    config.EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    if config.EXCEL_PATH.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = config.SHEET_NAME
    ws.append(headers())
    for cell in ws[1]:
        _style_header_cell(cell)
    ws.freeze_panes = "A2"
    wb.save(config.EXCEL_PATH)
    log.info("Nayi Excel file bani: %s", config.EXCEL_PATH)


def _sync_headers(ws) -> dict:
    """Sheet ke headers padhta hai; jo config ke fields sheet mein nahi hain unhe end mein jodta hai."""
    existing = {}
    for cell in ws[1]:
        if cell.value not in (None, ""):
            existing[str(cell.value)] = cell.column
    next_col = max(existing.values(), default=0) + 1
    for name in headers():
        if name not in existing:
            cell = ws.cell(row=1, column=next_col, value=name)
            _style_header_cell(cell)
            existing[name] = next_col
            next_col += 1
    return existing


def _last_data_row(ws) -> int:
    """Asli aakhri data wali row (khali formatted rows ignore karta hai)."""
    for row_no in range(ws.max_row, 1, -1):
        if any(c.value not in (None, "") for c in ws[row_no]):
            return row_no
    return 1


def _existing_ids(ws, col_map) -> set:
    col = col_map.get("Entry_ID")
    if not col:
        return set()
    ids = set()
    for row_no in range(2, ws.max_row + 1):
        value = ws.cell(row=row_no, column=col).value
        if value:
            ids.add(str(value))
    return ids


def _to_cell_value(field_type, value):
    if value is None:
        return None
    if field_type == "date":
        parsed = validator.parse_date(value)
        return parsed if parsed else value
    return value


def _autosize(ws):
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        longest = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(max(longest + 3, 12), 45)


def _backup():
    if not config.EXCEL_PATH.exists():
        return
    config.BACKUP_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = config.BACKUP_DIR / f"{config.EXCEL_PATH.stem}_{stamp}.xlsx"
    shutil.copy2(config.EXCEL_PATH, target)
    backups = sorted(config.BACKUP_DIR.glob(f"{config.EXCEL_PATH.stem}_*.xlsx"))
    for old in backups[:-config.MAX_BACKUPS]:
        old.unlink(missing_ok=True)


# ---------------------------------------------------------------- public API
def existing_ids() -> set:
    """UI preview ke liye: Excel mein pehle se kaun-kaun si entries hain."""
    if not config.EXCEL_PATH.exists():
        return set()
    try:
        wb = load_workbook(config.EXCEL_PATH, read_only=True)
        try:
            if config.SHEET_NAME not in wb.sheetnames:
                return set()
            ws = wb[config.SHEET_NAME]
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header or "Entry_ID" not in header:
                return set()
            idx = list(header).index("Entry_ID")
            return {str(r[idx]) for r in rows if len(r) > idx and r[idx]}
        finally:
            wb.close()
    except PermissionError:
        return set()


def append_records(records: list, source: str = "manual") -> dict:
    """
    Clean records ko Excel mein jodta hai.
    Returns: {"added": int, "duplicates": int}
    """
    try:
        _ensure_workbook()
        wb = load_workbook(config.EXCEL_PATH)
        ws = _get_sheet(wb)
        col_map = _sync_headers(ws)
        seen = _existing_ids(ws, col_map)
        row_no = _last_data_row(ws) + 1
        types = {f.name: f.type for f in config.FIELDS}
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")

        added, duplicates = 0, 0
        for record in records:
            entry_id = validator.make_entry_id(record)
            if entry_id in seen:
                duplicates += 1
                continue
            full = {**record, "Source": source, "Added_On": stamp, "Entry_ID": entry_id}
            for name, value in full.items():
                if name not in col_map:
                    continue
                cell = ws.cell(row=row_no, column=col_map[name])
                cell.value = _to_cell_value(types.get(name), value)
                cell.font = _BODY_FONT
                if types.get(name) == "date":
                    cell.number_format = "DD-MM-YYYY"
                elif types.get(name) == "number":
                    cell.number_format = "#,##0.00"
            seen.add(entry_id)
            row_no += 1
            added += 1

        if added:
            _backup()
            _autosize(ws)
            ws.freeze_panes = "A2"
            tmp = config.EXCEL_PATH.with_name(config.EXCEL_PATH.stem + ".tmp.xlsx")
            try:
                wb.save(tmp)
                os.replace(tmp, config.EXCEL_PATH)
            finally:
                tmp.unlink(missing_ok=True)
        log.info("Save: added=%s duplicates=%s source=%s", added, duplicates, source)
        return {"added": added, "duplicates": duplicates}
    except PermissionError as exc:
        log.warning("Excel locked: %s", exc)
        raise ExcelLockedError(
            f"'{config.EXCEL_PATH.name}' Excel mein khuli hai. Use band karke dobara try karo."
        ) from exc


def read_all() -> pd.DataFrame:
    """Records tab ke liye poora data DataFrame mein."""
    if not config.EXCEL_PATH.exists():
        return pd.DataFrame(columns=headers())
    try:
        return pd.read_excel(config.EXCEL_PATH, sheet_name=config.SHEET_NAME)
    except (ValueError, PermissionError):
        return pd.DataFrame(columns=headers())


def stats() -> dict:
    """Sidebar ke liye chhota summary."""
    df = read_all()
    number_fields = [f.name for f in config.FIELDS if f.type == "number"]
    label, total = "Total", 0.0
    if number_fields and number_fields[0] in df.columns:
        label = f"Total {number_fields[0]}"
        total = float(pd.to_numeric(df[number_fields[0]], errors="coerce").sum())
    return {"count": len(df), "total": total, "total_label": label}


if __name__ == "__main__":
    # Roadmap step 2: bina AI ke Excel mein dummy entry daal ke test karo
    #   python -m core.excel_store
    sample = {"Name": "Ramesh", "Amount": 5000.0, "Date": "24-09-2026",
              "Invoice_No": "INV-101", "Notes": "test entry"}
    print(append_records([sample], source="demo"))
    print(read_all().tail())
